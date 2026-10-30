"""Build the skill-based evaluation report (the college's Word template) as .docx and .pdf.

Follows the layout of the course's sample report page by page: cover, an empty certificates
page, title page, abstract, chapters 1-6, conclusion, references, screenshot appendix and the
evaluation sheet. The text describes this project, and every number in it is read at build time
from reports/tables/, reports/data_check.json, models/ or config.py; none is typed here.

  .venv/Scripts/python.exe scripts/build_skill_report.py           # reports/skill_report.docx + .pdf
  .venv/Scripts/python.exe scripts/build_skill_report.py --no-pdf  # .docx only

The PDF is printed by Microsoft Word (Windows, through PowerShell COM), so it looks exactly like
the .docx opened in Word. The UML figures come from scripts/report_diagrams.py.
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
import config  # noqa: E402
import report_diagrams  # noqa: E402

REPORTS = ROOT / "reports"
TABLES = REPORTS / "tables"
FIGURES = REPORTS / "figures"
SHOTS = REPORTS / "screenshots"
ASSETS = ROOT / "docs" / "report_assets"
OUT_DOCX = REPORTS / "skill_report.docx"
OUT_PDF = REPORTS / "skill_report.pdf"

# Front-matter details (the course's template; not results)
STUDENT = "ISAAC BENJAMIN DAVID"
REG_NO = "URK23AI1035"
COURSE_CODE = "25DC2005"
COURSE_NAME = "AutoML"
MONTH = "OCTOBER 2026"
TITLE = "Malware Detection on EMBER2024 with Manual Tuning vs AutoML"
TITLE_LINES = ["Malware Detection on EMBER2024 with", "Manual Tuning vs AutoML"]
REPO_URL = "https://github.com/Turbo-Soggy/automl"

BODY_FONT, COVER_FONT = "Times New Roman", "Arial"
LINK_BLUE = RGBColor(0x05, 0x63, 0xC1)
LINE_SPACING = 1.56      # matches the sample's measured 21.6 pt line pitch for Times New Roman 12 pt
HEADER_FILL = "DBE5F1"   # the light blue of the sample's evaluation-sheet header row


# --- Facts: every number the text uses, read from the saved results ----------------------------

def table(name: str) -> pd.DataFrame:
    return pd.read_csv(TABLES / f"{name}.csv")


def pct(x: float, digits: int = 1) -> str:
    return f"{100 * x:.{digits}f}%"


def facts() -> dict[str, Any]:
    """Formatted numbers for the text, keyed by name."""
    F: dict[str, Any] = {}
    dc = json.loads((REPORTS / "data_check.json").read_text(encoding="utf-8"))
    raw = {m.group(1): (int(m.group(2)), int(m.group(3)))
           for m in (re.match(r"(train|test): (\d+) of (\d+) raw lines", f) for f in dc["flags"]) if m}
    F["raw_lines"] = f"{sum(t for _, t in raw.values()):,}"
    F["dup_lines"] = f"{sum(d for d, _ in raw.values()):,}"
    cross = [int(m.group(1)) for m in (re.match(r"test: (\d+) sha256 appeared in two different weeks", f)
                                       for f in dc["flags"]) if m]
    F["cross_week"] = f"{sum(cross):,}"
    F["unique_train"] = f"{dc['cache']['shapes']['train']['meta_rows']:,}"
    F["unique_test"] = f"{dc['cache']['shapes']['test']['meta_rows']:,}"
    F["dim"] = f"{dc['cache']['feature_dim']:,}"
    F["n_families"] = len(dc["cache"]["feature_groups"])
    sizes = dc["splits"]["sizes"]
    F["n_train"], F["n_random"], F["n_future"] = (f"{sizes[k]:,}" for k in ("train", "test_random", "test_future"))
    F["train_weeks"], F["future_weeks"] = dc["dates"]["train"]["n_weeks"], dc["dates"]["test"]["n_weeks"]
    F["train_from"], F["train_to"] = dc["dates"]["train"]["date_min"][:10], dc["dates"]["train"]["date_max"][:10]
    F["future_from"], F["future_to"] = dc["dates"]["test"]["date_min"][:10], dc["dates"]["test"]["date_max"][:10]
    F["overlap"] = sum(dc["splits"]["sha256_overlap"].values())
    info = json.loads((ROOT / "data" / "splits" / "splits_info.json").read_text(encoding="utf-8"))
    F["tune_subset"] = f"{info['tune_subset']['n']:,}"
    F["trials"], F["folds"], F["seed"] = config.TUNE_TRIALS, config.CV_FOLDS, config.SEED
    F["budget_min"] = f"{config.AUTOML_BUDGET_S / 60:g}"
    F["n_boot"] = f"{config.N_BOOTSTRAP:,}"

    rvf = table("random_vs_future").set_index("model")
    F["rvf"] = rvf
    F["tuned_auc_r"], F["tuned_auc_f"] = f"{rvf.loc['optuna', 'test_random_roc_auc']:.4f}", f"{rvf.loc['optuna', 'test_future_roc_auc']:.4f}"
    F["tuned_gap"] = f"{rvf.loc['optuna', 'roc_auc_gap']:.4f}"
    F["tuned_gap_ci"] = f"[{rvf.loc['optuna', 'roc_auc_gap_ci_low']:.4f}, {rvf.loc['optuna', 'roc_auc_gap_ci_high']:.4f}]"
    F["tuned_tpr_f"] = pct(rvf.loc["optuna", "test_future_tpr_at_1pct_fpr"])
    F["default_tpr_f"] = pct(rvf.loc["lgbm_default", "test_future_tpr_at_1pct_fpr"])
    F["logreg_gap"] = f"{rvf.loc['logreg', 'roc_auc_gap']:.4f}"
    F["logreg_auc_f"] = f"{rvf.loc['logreg', 'test_future_roc_auc']:.4f}"
    trees = ["optuna", "grid", "random", "flaml", "autogluon", "h2o", "lgbm_default"]
    F["tree_gap_min"], F["tree_gap_max"] = (f"{rvf.loc[trees, 'roc_auc_gap'].min():.4f}",
                                            f"{rvf.loc[trees, 'roc_auc_gap'].max():.4f}")
    hand = rvf.loc[["grid", "random", "optuna"], "test_future_tpr_at_1pct_fpr"]
    F["hand_spread_pts"] = f"{100 * (hand.max() - hand.min()):.2f}"

    avm = table("automl_vs_manual").set_index("method")
    F["avm"] = avm
    F["ag_min"], F["tuned_min"] = f"{avm.loc['autogluon', 'wall_minutes']:.1f}", f"{avm.loc['lgbm_tuned', 'wall_minutes']:.1f}"
    F["all_manual_min"] = f"{avm.loc['lgbm_tuned', 'all_manual_seconds'] / 60:.1f}"
    pw = table("automl_pairwise")
    fut = pw[(pw.test_set == "test_future") & (pw.method_b == "lgbm_tuned")]

    def diff(tool: str, metric: str) -> pd.Series:
        return fut[(fut.method_a == tool) & (fut.metric == metric)].iloc[0]
    F["pw"] = {t: {m: diff(t, m) for m in ("roc_auc", "tpr_at_1pct_fpr")} for t in ("flaml", "autogluon", "h2o")}
    F["ag_tied"] = all(not F["pw"]["autogluon"][m].significant_95 for m in ("roc_auc", "tpr_at_1pct_fpr"))

    ts = table("tuning_summary").set_index("method")
    F["ts"] = ts
    F["best_cv_method"] = ts.loc[["grid", "random", "optuna"], "best_cv_auc"].idxmax()
    thr = json.loads((ROOT / "models" / "threshold.json").read_text(encoding="utf-8"))
    F["thr"] = f"{thr['threshold']:.4f}"
    F["cost_ratio"] = f"{thr['cost_ratio_fn_to_fp']:g}"
    F["dep_recall"], F["dep_fpr"] = pct(thr["test_future_check"]["recall"]), pct(thr["test_future_check"]["fpr"])

    prev = table("precision_both_sets")
    p = prev[(prev.model == "optuna") & (prev.test_set == "test_future") & (prev.threshold_name == "deployment")]
    p = p[p.prevalence == p.prevalence.min()].iloc[0]
    F["prev_rate"], F["prec_mix"], F["prec_low"] = pct(p.prevalence, 0), pct(p.precision_at_test_mix, 1), pct(p.precision, 1)
    F["false_per_true"] = f"{p.false_alarms_per_10k / p.true_alarms_per_10k:.1f}"

    leak = table("leakage")
    dr = leak[leak.variant == "leak_detection_ratio"]
    tools = dr[dr.model != "lgbm_tuned"]
    F["leak_relied"], F["leak_models"] = int(dr["relied_on_leak"].sum()), len(dr)
    F["leak_rank_max"] = int(dr["leak_rank_shap"].max())
    F["leak_kept"] = int((tools["leak_dropped_by_tool"] == False).sum())  # noqa: E712 (NaN for the manual model)
    F["leak_silent"] = int((tools["n_lines_mentioning_leak"] == 0).sum())
    F["leak_tools"] = len(tools)
    F["leak_budget_min"] = f"{tools['budget_s'].iloc[0] / 60:g}"
    F["no_tool_caught"] = F["leak_kept"] == F["leak_tools"] and F["leak_silent"] == F["leak_tools"]
    F["leak_verdict"] = ("no AutoML tool removed it or mentioned it in its logs" if F["no_tool_caught"] else
                         f"{F['leak_tools'] - F['leak_kept']} of {F['leak_tools']} AutoML tools removed it")
    F["leak_auc_min"], F["leak_auc_max"] = (f"{dr['test_future_roc_auc'].min():.4f}",
                                            f"{dr['test_future_roc_auc'].max():.4f}")
    F["leak_rank_text"] = ("their most important input" if F["leak_rank_max"] == 1
                           else f"a top-{F['leak_rank_max']} input")
    fs = leak[leak.variant == "leak_first_submission"]
    F["first_sub_max_delta"] = f"{fs['delta_test_future_roc_auc'].abs().max():.4f}"
    ab = leak[leak.variant == "ablate_time_proxies"].iloc[0]
    F["ablate_delta"] = f"{ab.delta_test_future_roc_auc:+.4f}"

    ch = table("challenge")
    c = ch[(ch.model == "optuna") & (ch.threshold_name == "deployment") & (ch.subset == "all Win64 challenge")].iloc[0]
    F["chal_n"], F["chal_rate"], F["chal_vs"] = int(c.n), pct(c.detection_rate), pct(c.test_future_tpr_same_threshold)

    bias = table("bias_slices")
    sg = bias[(bias.test_set == "test_future") & (bias.slice_family == "signed")
              & (bias.threshold_name == "train_oof_1pct_fpr")].set_index("slice")
    F["signed_tpr"], F["unsigned_tpr"] = pct(sg.loc["signed", "tpr"]), pct(sg.loc["unsigned", "tpr"])
    F["signed_n"] = f"{int(sg.loc['signed', 'n_malicious']):,}"

    cw = table("class_weight")
    bal = cw[(cw.test_set == "test_future") & (cw.model == "class_weight_balanced")].iloc[0]
    F["cw_delta"] = f"{100 * bal.tpr_at_1pct_fpr_minus_baseline:+.2f}"
    F["cw_ci"] = (f"[{100 * bal.tpr_at_1pct_fpr_minus_baseline_ci_low:+.2f}, "
                  f"{100 * bal.tpr_at_1pct_fpr_minus_baseline_ci_high:+.2f}]")

    cv = table("cv_comparison").set_index("estimate")
    F["cv_strat_auc"] = f"{cv.loc['stratified_3fold_cv', 'roc_auc']:.4f}"
    F["cv_fwd_auc"] = f"{cv.loc['forward_cv_4fold', 'roc_auc']:.4f}"
    F["cv_strat_err"] = f"{cv.loc['stratified_3fold_cv', 'tpr_at_1pct_fpr_abs_error_vs_future'] * 100:.2f}"
    F["cv_fwd_err"] = f"{cv.loc['forward_cv_4fold', 'tpr_at_1pct_fpr_abs_error_vs_future'] * 100:.2f}"

    cal = table("calibration").set_index("model")
    F["brier_tuned"] = f"{cal.loc['lgbm_tuned', 'brier']:.4f}"
    F["brier_logreg"] = f"{cal.loc['logreg', 'brier']:.4f}"

    fam = table("shap_family_importance").sort_values("rank")
    F["top_families"] = ", ".join(f"{r.family} ({pct(r.share, 0)})" for _, r in fam.head(3).iterrows())

    drift = table("drift")
    allw = drift[drift.week.astype(str) == "all_test_future"].set_index("model")
    F["drift_exp"], F["drift_slide"] = (f"{allw.loc['expanding', 'roc_auc']:.4f}", f"{allw.loc['sliding_12w', 'roc_auc']:.4f}")

    schema = json.loads((ROOT / "models" / "final" / "schema.json").read_text(encoding="utf-8"))
    F["num_trees"] = schema["num_trees"]
    F["model_mb"] = f"{(ROOT / 'models' / 'final' / schema['model_file']).stat().st_size / 2**20:.1f}"
    return F


def behind_phrase(F: dict[str, Any]) -> str:
    """How FLAML and H2O compare with tuned LightGBM on the main metric, from the paired bootstrap."""
    rows = [F["pw"][t]["tpr_at_1pct_fpr"] for t in ("flaml", "h2o")]
    if all(r.significant_95 and r.diff_a_minus_b < 0 for r in rows):
        return "slightly but significantly behind tuned LightGBM"
    return "not clearly different from tuned LightGBM"


def leak_tools_sentence(F: dict[str, Any]) -> str:
    """What the AutoML tools did with the leaky column, from the leakage table."""
    n, kept, silent = F["leak_tools"], F["leak_kept"], F["leak_silent"]
    if kept == n and silent == n:
        return f"All {n} AutoML tools kept the column and none mentioned it in their logs."
    return f"Of the {n} AutoML tools, {kept} kept the column and {silent} never mentioned it in their logs."


def plain_pick(note: str) -> str:
    """Plain words for an AutoML tool's final model, read from its interpretability note."""
    if "xgb" in note:
        return "XGBoost, a boosted-tree model"
    if "all weight on LightGBM" in note:
        return "an ensemble that put all its weight on one LightGBM model"
    if "weight only to GBM" in note:
        return "an ensemble that in effect used one boosted-tree model (GBM)"
    return note


MODEL_NAMES = {"optuna": "LightGBM, Optuna (tuned)", "grid": "LightGBM, grid search", "random": "LightGBM, random search",
               "lgbm_default": "LightGBM, default settings", "logreg": "Logistic regression", "flaml": "FLAML",
               "autogluon": "AutoGluon", "h2o": "H2O AutoML"}
TOOL_NAMES = {"flaml": "FLAML", "autogluon": "AutoGluon", "h2o": "H2O AutoML"}


# --- Word helpers ------------------------------------------------------------------------------

def setup(doc: Document) -> None:
    """A4, 2.54 cm side and top margins, Times New Roman 12 pt, 1.5 lines, justified: the sample's body style."""
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin"):
        setattr(s, side, Cm(2.54))
    s.bottom_margin = Cm(1.0)   # the sample's text runs to about 0.9 cm above the page edge
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = BODY_FONT, Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing, pf.space_after, pf.space_before = LINE_SPACING, Pt(0), Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def runs(p: Paragraph, text: str, font: str | None = None, size: float | None = None,
         bold: bool = False, italic: bool = False) -> Paragraph:
    """Add text with **bold** and *italic* markup to a paragraph."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        r = p.add_run(part.strip("*") if part.startswith("*") else part)
        r.bold = bold or part.startswith("**")
        r.italic = italic or (part.startswith("*") and not part.startswith("**"))
        if font:
            r.font.name = font
            r.element.rPr.rFonts.set(qn("w:eastAsia"), font)
        if size:
            r.font.size = Pt(size)
    return p


def para(doc: Document, text: str = "", *, align: str = "justify", size: float | None = None, bold: bool = False,
         italic: bool = False, font: str | None = None, before: float = 0, after: float = 0,
         spacing: float | None = None, indent: float = 0, hanging: float = 0, keep: bool = False) -> Paragraph:
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if spacing is not None:
        pf.line_spacing = spacing
    if indent or hanging:
        pf.left_indent, pf.first_line_indent = Cm(indent), Cm(-hanging)
    pf.keep_with_next = keep
    return runs(p, text, font, size, bold, italic)


def chapter(doc: Document, n: int, title: str, big_title: bool) -> None:
    """CHAPTER n / TITLE, as in the sample: chapters 1-2 use a larger title right below, 3-6 a gap first."""
    page_break(doc, gap=19)
    para(doc, f"CHAPTER {n}", align="center", size=16, bold=True)
    if big_title:
        para(doc, title, align="center", size=16, bold=True, after=10)
    else:
        para(doc, title, align="center", size=14, bold=True, before=14, after=8)


def heading(doc: Document, text: str, before: float = 0) -> None:
    para(doc, text, align="left", bold=True, before=before, keep=True)


def bullets(doc: Document, items: list[str], level: int = 0, numbered: bool = False) -> None:
    """Bullets (• then o) or a numbered list, with a hanging indent like the sample."""
    for i, item in enumerate(items, start=1):
        mark = f"{i}." if numbered else ("•" if level == 0 else "o")
        p = para(doc, "", indent=1.27 + 1.27 * level, hanging=0.63)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.27 + 1.27 * level))
        runs(p, f"{mark}\t{item}")


def page_break(doc: Document, gap: float = 0) -> None:
    """Start a new page, then leave `gap` points above the first line.

    Word drops "space before" on the first paragraph after a manual page break, so the gap is an empty
    paragraph of exact height instead; the paragraph that holds the break is kept 1 pt high.
    """
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    if gap:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(gap)


def picture(doc: Document, path: Path, width_cm: float, before: float = 6, after: float = 6) -> None:
    p = para(doc, align="center", before=before, after=after, spacing=1.0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def caption(doc: Document, text: str) -> None:
    para(doc, text, align="center", bold=True, before=4, after=10)


def shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def data_table(doc: Document, header: list[str], rows: list[list[str]], widths: list[float],
               size: float = 10, center_from: int = 1) -> None:
    """A grid table with a shaded bold header, compact single-spaced cells."""
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, values in enumerate([header] + rows):
        for c, value in enumerate(values):
            cell = t.cell(r, c)
            cell.width = Cm(widths[c])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c >= center_from or r == 0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
            runs(p, value, size=size, bold=(r == 0))
            if r == 0:
                shade(cell, HEADER_FILL)


def hyperlink(p: Paragraph, url: str, text: str | None = None) -> None:
    """A clickable blue underlined link (python-docx has no API for this)."""
    rel = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), str(LINK_BLUE))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text or url
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    p._p.append(link)


def page_number_footer(section: Any) -> None:
    """Centred PAGE field in this section's footer, restarting at 1 (the sample numbers only its last page)."""
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    section._sectPr.append(pg)


# --- Front matter -------------------------------------------------------------------------------

def cover(doc: Document) -> None:
    picture(doc, ASSETS / "karunya_banner.png", 15.7, before=0, after=0)
    for text, size, before in (("DIVISION OF DATA SCIENCE AND CYBER SECURITY", 14, 60),
                               ("SCHOOL OF COMPUTER SCIENCE AND TECHNOLOGY", 14, 32),
                               ("A SKILL BASED EVALUATION REPORT", 17, 50)):
        para(doc, text, align="center", size=size, bold=True, font=COVER_FONT, before=before, spacing=1.15)
    para(doc, TITLE_LINES[0], align="center", size=16, bold=True, font=COVER_FONT, before=18, spacing=1.15)
    para(doc, TITLE_LINES[1], align="center", size=16, bold=True, font=COVER_FONT, spacing=1.15)
    para(doc, "SUBMITTED BY", align="center", size=14, bold=True, font=COVER_FONT, before=33, spacing=1.15)
    para(doc, f"{STUDENT}({REG_NO})", align="center", size=14, bold=True, font=COVER_FONT, before=8, spacing=1.15)
    para(doc, "COURSE CODE", align="center", size=14, bold=True, font=COVER_FONT, before=57, spacing=1.15)
    para(doc, COURSE_CODE, align="center", size=14, bold=True, font=COVER_FONT, spacing=1.15)
    para(doc, "COURSE NAME", align="center", size=14, bold=True, font=COVER_FONT, before=56, spacing=1.15)
    para(doc, COURSE_NAME, align="center", size=14, bold=True, font=COVER_FONT, spacing=1.15)
    para(doc, MONTH, align="center", size=14, bold=True, font=COVER_FONT, before=30, spacing=1.15)


def certificates(doc: Document) -> None:
    """Left empty on purpose: the student adds the certificates."""
    page_break(doc, gap=30)
    para(doc, "CERTIFICATES", align="center", size=14, bold=True, font=COVER_FONT)


def title_page(doc: Document) -> None:
    page_break(doc, gap=55)
    para(doc, "TITLE", align="center", size=16, bold=True, spacing=1.15)
    para(doc, TITLE_LINES[0], align="center", size=16, bold=True, font=COVER_FONT, before=20, spacing=1.15)
    para(doc, TITLE_LINES[1], align="center", size=16, bold=True, font=COVER_FONT, spacing=1.15)
    para(doc, "Submitted by", align="center", size=14, bold=True, italic=True, before=66, spacing=1.15)
    para(doc, f"{STUDENT}({REG_NO})", align="center", size=16, bold=True, before=8, spacing=1.15)
    picture(doc, ASSETS / "karunya_emblem.png", 3.5, before=138, after=0)
    para(doc, "DIVISION OF DATA SCIENCE AND CYBER SECURITY", align="center", size=15, bold=True, before=50,
         spacing=1.15)
    para(doc, "KARUNYA INSTITUTE OF TECHNOLOGY AND SCIENCES", align="center", size=14, bold=True, before=56,
         spacing=1.0)
    para(doc, "(Declared as Deemed-to-be-under Sec-3 of the UGC Act,", align="center", size=16, bold=True,
         spacing=1.0)
    para(doc, "1956) Karunya Nagar, Coimbatore - 641 114. INDIA", align="center", size=16, bold=True, spacing=1.0)
    para(doc, f"{MONTH}.", align="center", size=16, bold=True, before=30, spacing=1.0)


def abstract(doc: Document, F: dict[str, Any]) -> None:
    page_break(doc, gap=21)
    para(doc, "ABSTRACT", align="center", size=14, bold=True, after=22)
    pw = F["pw"]["autogluon"]["tpr_at_1pct_fpr"]
    for text in [
        "This project presents a comparative study of manual hyperparameter tuning and Automated Machine Learning "
        "(AutoML) for static malware detection. The task is to label Windows 64-bit Portable Executable (PE) files as "
        "benign or malicious from features extracted without running them. The project uses the EMBER2024 dataset "
        f"(its Win64 part), in which every file is described by {F['dim']} numeric features in {F['n_families']} "
        "families. It also asks a research question: how much does a random train/test split overstate real-world "
        "performance, and do AutoML tools protect against data leakage and temporal drift, or fall for them?",
        f"After removing duplicate files, a stratified sample of {F['n_train']} training files and a random test set of "
        f"{F['n_random']} files were drawn from weeks 0-{F['train_weeks'] - 1}, and a future test set of "
        f"{F['n_future']} files from the {F['future_weeks']} later weeks, with a fixed random seed of {F['seed']}. "
        "LightGBM with default settings and a logistic-regression model serve as baselines. LightGBM was then tuned "
        f"with Grid Search, Random Search and Optuna, each with {F['trials']} configurations and {F['folds']}-fold "
        f"cross-validation. FLAML, AutoGluon and H2O AutoML were each given a {F['budget_min']}-minute budget on the "
        "same training data.",
        "The main metric is the share of malware caught when only 1% of benign files may be flagged (TPR at 1% FPR), "
        f"with ROC-AUC, F1, accuracy and time, and 95% bootstrap confidence intervals. On the future test set, tuned "
        f"LightGBM caught {F['tuned_tpr_f']} of malware; AutoGluon was statistically tied with it "
        f"({100 * pw.diff_a_minus_b:+.2f} points, 95% CI {100 * pw.ci_low:+.2f} to {100 * pw.ci_high:+.2f}) in "
        f"{F['ag_min']} minutes against {F['tuned_min']} for the selected manual search, while FLAML and H2O AutoML "
        f"were {behind_phrase(F)}. The random split overstated the tuned model's AUC by only {F['tuned_gap']}, but "
        f"a planted leaky feature was used as the top input by {F['leak_relied']} of {F['leak_models']} models and "
        f"{F['leak_verdict']}.",
        "SHAP explains the tuned model's decisions. The tuned LightGBM was exported with a cost-based threshold and "
        "deployed through a FastAPI service and a Streamlit presentation page, with a model card. The study shows "
        "where AutoML matches careful manual work, and where human judgement is still needed: choosing honest test "
        "splits, spotting leaks, and setting thresholds for a realistic share of malware.",
    ]:
        para(doc, text)


# --- Chapters -----------------------------------------------------------------------------------

def introduction(doc: Document, F: dict[str, Any]) -> None:
    chapter(doc, 1, "INTRODUCTION", big_title=True)
    for text in [
        "Malware detection is one of the main uses of machine learning in security. New malicious programs appear "
        "every day, and a detector that only matches known signatures misses them. Static detection reads the "
        "structure of a file (its headers, sections, imported functions, embedded text and digital signature) "
        "without running it, and a model learns which patterns separate malware from benign software. Building such "
        "a model involves data preparation, model selection, hyperparameter tuning, evaluation and deployment, "
        "which take expertise and time. Automated Machine Learning (AutoML) aims to automate model selection and "
        "tuning.",
        f"This project, titled **\"{TITLE},\"** compares the two approaches on the EMBER2024 benchmark, a public "
        f"dataset of Windows programs labelled as malware or benign. Each file is described by {F['dim']} numbers "
        "produced by the dataset's feature extractor (thrember). The prediction task is binary classification.",
        "The primary objective is to evaluate whether AutoML frameworks (FLAML, AutoGluon and H2O AutoML) can match "
        "or beat a carefully tuned LightGBM model, the standard model for EMBER. LightGBM is tuned with Grid Search, "
        "Random Search and Optuna under an equal budget, and the AutoML tools receive the same training data and an "
        "equal time budget each.",
        "A second objective is to test how honest the usual evaluation is. Every model is scored on two test sets: a "
        "random test set from the same weeks as the training data, and a future test set from later weeks, which is "
        "closer to real use. The project also plants leaky features on purpose to see whether any tool notices them, "
        "measures performance on slices of the data, explains the model with SHAP, and checks calibration and drift.",
        "The main metric is the share of malware caught at a 1% false-positive rate, because a security team can only "
        "handle a limited number of false alarms. Cross-validation on the training data is used for every choice; the "
        "two test sets are used only for the final scores. The chosen model is served through a FastAPI service and a "
        "Streamlit presentation page.",
        "Overall, the project aims to give a practical understanding of the strengths and limits of manual tuning and "
        "AutoML, and of the evaluation practices that decide whether a malware detector can be trusted.",
    ]:
        para(doc, text)


def literature(doc: Document) -> None:
    chapter(doc, 2, "LITERATURE REVIEW", big_title=True)
    para(doc, "Machine learning has been widely applied to malware detection from static features. Traditional "
              "approaches require the developer to choose an algorithm, tune its hyperparameters and design an "
              "evaluation. These steps give control over the modelling process but take extensive experimentation. "
              "Automated Machine Learning reduces this effort by automating model selection and tuning.", before=8)
    sections = [
        ("Static Malware Detection and the EMBER Datasets",
         "EMBER (Anderson and Roth, 2018) released features of over a million Windows PE files with a LightGBM "
         "baseline, and became the standard benchmark for static detection. EMBER2024 (Joyce et al., 2025) updates "
         "it with files from September 2023 to December 2024, weekly releases, several file types and a "
         "\"challenge set\" of malware that antivirus engines first missed. Its feature extractor, thrember, turns "
         "each file into a fixed-length vector of byte statistics, header fields, section, import and export "
         "information, strings and signature details."),
        ("Manual Machine Learning and Hyperparameter Optimization",
         "Gradient-boosted decision trees are strong on tabular data. LightGBM (Ke et al., 2017) builds trees "
         "leaf-wise with histogram-based splits, which makes it fast on wide feature sets such as EMBER. Grid Search "
         "evaluates a fixed set of combinations, Random Search samples combinations at random and can reach "
         "similar results with fewer trials, and Optuna (Akiba et al., 2019) uses the Tree-structured Parzen Estimator to propose "
         "promising configurations from earlier trials. Scikit-learn (Pedregosa et al., 2011) provides the "
         "cross-validation, logistic regression and metric tools."),
        ("Automated Machine Learning Frameworks",
         "FLAML (Wang et al., 2021) searches learners and hyperparameters while accounting for training cost. "
         "AutoGluon-Tabular (Erickson et al., 2020) trains many model types with presets and combines them by "
         "stacking and weighted ensembling. H2O AutoML (LeDell and Poirier, 2020) trains a sequence of algorithms "
         "and builds stacked ensembles of the best of each family. All three reduce manual work, but their results "
         "depend on the time budget, the hardware and how they validate internally."),
        ("Model Evaluation and Explainability",
         "Reliable evaluation is essential when comparing approaches. For malware detection the true-positive rate at "
         "a low false-positive rate matters more than accuracy, and precision depends strongly on how common malware "
         "is. Paired bootstrap confidence intervals show whether a difference between two models is larger than the "
         "noise of the test set. SHAP (Lundberg and Lee, 2017), and its exact TreeSHAP form for tree models, explains "
         "how each feature contributed to a prediction; such explanations help understanding but do not prove a "
         "model correct. Model cards (Mitchell et al., 2019) document intended use, metrics and limits."),
    ]
    for title, text in sections:
        heading(doc, title)
        para(doc, text)
    heading(doc, "Temporal Drift and Data Leakage")
    para(doc, "Malware changes over time, so a random split, which mixes old and new files in both training and "
              "test data, can overstate real performance. TESSERACT (Pendlebury et al., 2019) showed that "
              "time-ignorant evaluation inflates malware-detection results, and recommended testing on files that are "
              "strictly later than the training data. Data leakage, where a feature carries information about the "
              "label that would not exist in real use, can produce misleadingly high scores; deliberately adding a "
              "leaky feature tests whether a pipeline, or an AutoML tool, notices it.")
    para(doc, "In summary, existing libraries and AutoML frameworks offer different ways to build a detector. Manual "
              "methods give more direct control, whereas AutoML can reduce experimentation time and explore many "
              "configurations. This project brings both together on the same malware-detection task, and evaluates "
              "them on both a random and a future test set.")


def system_design(doc: Document, F: dict[str, Any]) -> None:
    chapter(doc, 3, "SYSTEM DESIGN", big_title=False)
    heading(doc, "3.1 Architecture and high-level design", before=10)
    picture(doc, FIGURES / "architecture.png", 15.9)
    caption(doc, "Fig 3.1.1 System diagram")
    para(doc, "The system has four layers.")
    bullets(doc, [
        f"Data layer. The EMBER2024 Win64 files ({F['raw_lines']} raw records in weekly files) are downloaded from "
        "Hugging Face with checksum checks. src/data.py removes duplicate files by sha256, extracts the feature "
        f"vectors with thrember into a cache, and makes the fixed splits (train {F['n_train']}, test_random "
        f"{F['n_random']}, test_future {F['n_future']}) with a random seed of {F['seed']}.",
        "Machine learning layer (src/). It has four stages:",
    ], numbered=True)
    bullets(doc, [
        "Step 2 builds the LightGBM and logistic-regression baselines and tunes LightGBM with Grid Search, Random "
        "Search and Optuna (src/tune.py).",
        "Step 3 runs FLAML, AutoGluon and H2O AutoML with the same time budget (src/automl_runs.py).",
        "Step 4 does the critical evaluation: random vs future split, leakage experiments, bias slices, SHAP, "
        "calibration, a cost-based threshold and drift.",
        "Step 5 exports the selected model with its schema and threshold (scripts/export_model.py).",
    ], level=1)
    p = para(doc, "", indent=1.27, hanging=0.63)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(1.27))
    runs(p, "3.\tArtifact layer. The deployed model is saved in models/final/ (model.txt, schema.json with its sha256 "
            "and threshold). Every run appends a record to reports/experiments.jsonl, and scripts turn it into "
            "tables, figures, the model card and this report.")
    p = para(doc, "", indent=1.27, hanging=0.63)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(1.27))
    runs(p, "4.\tServing layer (app/). A FastAPI backend exposes the model as a web service, and a Streamlit page "
            "presents the project and calls the API for live predictions. Both use the same saved model.")
    para(doc, "Design decisions.")
    bullets(doc, [
        "Tuning and model selection use cross-validation on the training data only; both test sets stay unseen.",
        "The pipeline is modular: each step is a separate script that reads and writes files, so any step can be "
        "rerun alone (run_all.py runs them in order).",
        "One configuration file (config.py) holds the seed, sizes, budgets and paths.",
        "Every number in the reports is generated from the experiment log and saved predictions, never typed by hand.",
    ])

    heading(doc, "3.2 UML diagrams", before=8)
    heading(doc, "Class diagram")
    picture(doc, report_diagrams.CLASS_PNG, 13.5)
    para(doc, "The project is mostly written as Python modules and functions, so the diagram shows each module as a "
              "component class. The configuration and data modules feed the tuning, AutoML and evaluation steps; "
              "every step logs to the experiment log. The export step saves the tuned LightGBM, which the "
              "Predictor class of the API loads and serves. The Streamlit page calls the API, or uses the same "
              "Predictor in-process when no API server is running.")
    heading(doc, "Sequence diagram", before=8)
    picture(doc, report_diagrams.SEQUENCE_PNG, 13.5)
    para(doc, "A user opens the page, which reads the saved result tables for the presentation sections, and picks an "
              f"example file. The page sends its {F['dim']} features to POST /predict. The API checks that the vector "
              "has the right length and only finite values, scores it with the LightGBM model, computes TreeSHAP "
              "values, applies the threshold and sums the SHAP values by feature family. The page shows the verdict, "
              "the flag line (the score above which a file is flagged) and the four families that pushed the "
              "decision most. As an alternative path, an "
              "uploaded program goes to POST /predict_file, where it is parsed in memory and never saved or run.")

    heading(doc, "3.3 User interface design (Streamlit)", before=8)
    para(doc, "The interface is a single web page titled \"Can a model spot malware without running it?\" that reads "
              "top to bottom as the whole presentation, in nine numbered sections, with no tabs or side panels.")
    bullets(doc, [
        "**Story sections.** The question, the data (with a timeline of the training and future weeks), the "
        "contenders, how they were compared, the results (a chart of every model with 95% confidence intervals), "
        "the winner and why, what was learned, and the limits. Every number is read from reports/tables/.",
        "**Try it live.** Five buttons pick a caught malware file, a safe program, a missed malware file, a false "
        "alarm, or an upload of a Windows program. The page shows a red or green verdict card with the malware "
        "score, the flag line and the true answer, then the feature families that pushed the decision, in plain "
        "words.",
        "**Technical detail.** A box at the end of the demo shows the raw probability, the threshold, the file's "
        "sha256 and the top-10 SHAP columns, for questions.",
    ])
    para(doc, "The FastAPI backend has its own interface at /docs (Swagger UI), where requests can be tried and the "
              "JSON response seen directly. present.py starts the API and the page together.")


def implementation(doc: Document, F: dict[str, Any]) -> None:
    chapter(doc, 4, "IMPLEMENTATION", big_title=False)
    para(doc, f"The implementation of the project, *{TITLE}*, focuses on building and comparing manually tuned and "
              "AutoML models for detecting malicious Windows programs. It is written in Python 3.12, and every step "
              "is a script that can be rerun on its own.")
    items = [
        ("1. Dataset Loading and Preprocessing",
         f"The EMBER2024 Win64 files are downloaded from Hugging Face and checked. Each weekly file repeats every "
         f"sample, so {F['dup_lines']} of the {F['raw_lines']} raw lines were duplicates within a file, and "
         f"{F['cross_week']} more files appeared in two different weeks; all were removed by sha256, keeping the "
         f"first occurrence and "
         f"leaving {F['unique_train']} files in the training period and {F['unique_test']} in the test period. "
         f"thrember turns each file into {F['dim']} float32 features, saved in a memory-mapped cache with the label, "
         f"week and metadata. The splits are stratified by label and week: train {F['n_train']} and test_random "
         f"{F['n_random']} from weeks 0-{F['train_weeks'] - 1}, and test_future {F['n_future']} from weeks "
         f"{F['train_weeks']}-{F['train_weeks'] + F['future_weeks'] - 1}, all with seed {F['seed']}. The leaky "
         "detection_ratio and first_submission_date fields are kept in the metadata only, never in the features."),
        ("2. Manual Machine Learning Implementation",
         "LightGBM with default settings and a logistic regression on standardized features are the baselines. "
         f"LightGBM is tuned with Grid Search, Random Search and Optuna (TPE sampler), each evaluating "
         f"{F['trials']} configurations with the same {F['folds']}-fold cross-validation folds on a {F['tune_subset']}"
         "-row subset of the training data, with early stopping inside each fold. The configuration with the best "
         "mean cross-validated ROC-AUC is refit on all training rows with a fixed number of rounds."),
        ("3. AutoML Implementation",
         f"FLAML, AutoGluon (medium-quality preset) and H2O AutoML each receive the same {F['n_train']} training rows "
         f"and a {F['budget_min']}-minute budget, with ROC-AUC as the metric. Seed {F['seed']} is passed to each "
         "tool's main seed argument, although some internal models keep the library's own fixed seed. Each tool chooses its own internal validation inside the training data. The tools run one "
         "at a time on the same machine, and a heartbeat guard records whether the machine slept, which would "
         "invalidate a timed run. Each tool's leaderboard, preprocessing notes and feature importance are saved."),
        ("4. Model Evaluation and Explainability",
         "Every model is scored on both test sets with ROC-AUC, F1, accuracy, false-positive rate and TPR at 1% FPR, "
         f"with 95% confidence intervals from {F['n_boot']} paired bootstrap resamples. The critical evaluation "
         "compares stratified and forward (week-blocked) cross-validation, reruns the tuned model and every AutoML "
         f"tool ({F['leak_budget_min']}-minute budget) with and without each leaky feature, measures recall and "
         "false alarms per slice (week, file size, packer, signed or unsigned), tries class weights, explains the "
         "model with TreeSHAP, checks calibration, chooses a cost-based threshold and tracks drift over the future "
         "weeks."),
        ("5. Application Development",
         f"The tuned LightGBM ({F['num_trees']} trees) is exported to models/final/ with schema.json, which records "
         f"its sha256, the feature layout and the threshold ({F['thr']}, chosen so that a missed malware costs "
         f"{F['cost_ratio']} times a false alarm, on training out-of-fold scores only). This threshold is "
         f"deliberately low: on the future test set it catches {F['dep_recall']} of malware but flags "
         f"{F['dep_fpr']} of benign files. A FastAPI service serves "
         "/health, /predict, /predict_batch and /predict_file, and a Streamlit page presents the project and the "
         "live demo."),
    ]
    for title, text in items:
        heading(doc, title)
        para(doc, text)


def testing(doc: Document, F: dict[str, Any]) -> None:
    chapter(doc, 5, "TESTING AND VALIDATION", big_title=False)
    para(doc, f"Testing and validation are essential parts of the project, *{TITLE}*. They verify that the data and "
              "code are correct, that no test information leaks into model choices, and that the application "
              "returns the same predictions as the offline model.")
    items = [
        ("1. Data Validation",
         "scripts/verify_data.py writes reports/data_check.json. It checks that every downloaded file has the "
         f"expected number of lines, that labels are only 0 or 1, that the {F['n_families']} feature groups add up to "
         f"{F['dim']} columns, that features and metadata stay aligned, and that the splits do not overlap "
         f"({F['overlap']} shared sha256 across them). Every future-test file is dated after every training file: "
         f"training covers {F['train_from']} to {F['train_to']} and the future test {F['future_from']} to "
         f"{F['future_to']}."),
        ("2. Model Validation",
         "The training set is used for model development and hyperparameter tuning; both test sets are reserved for "
         "final evaluation. All tuning methods share the same saved cross-validation folds. A week-blocked forward "
         f"cross-validation was compared with the stratified one: stratified CV gave AUC {F['cv_strat_auc']} and "
         f"forward CV {F['cv_fwd_auc']}, and their TPR at 1% FPR missed the future-test value by "
         f"{F['cv_strat_err']} and {F['cv_fwd_err']} points."),
        ("3. Performance Testing",
         "All eight models (two baselines, three tuned LightGBM searches and three AutoML tools) are evaluated with "
         "the same metric code (src/eval.py). The primary metric is TPR at 1% FPR, supported by ROC-AUC, F1, "
         "accuracy and false-positive rate. Paired bootstrap confidence intervals decide whether two models really "
         "differ."),
        ("4. Overfitting and Data Leakage Testing",
         f"The gap between the random and future test sets shows how much a random split overstates performance: "
         f"{F['tuned_gap']} AUC for tuned LightGBM, 95% CI {F['tuned_gap_ci']}. In the leakage experiment, the "
         "detection_ratio field (the share of antivirus engines that flag the file) is added as a feature, then removed; a "
         "second, subtler variant adds the first-submission date. The tuned model is also retrained without the "
         "time-proxy columns (compile and signing times) to check that they are not a hidden leak."),
        ("5. Application and Functional Testing",
         "Pytest covers split disjointness and temporal order, the tuners and metrics on synthetic data, the AutoML "
         "and leakage table code, and the application. The API tests check that a wrong-length or non-finite vector "
         "is rejected with HTTP 422, that an uploaded file is never written to disk, and that the API returns exactly "
         "the offline model's scores. The page tests render every section and compare the numbers shown with the "
         "result tables."),
        ("6. Validation Summary",
         "The final validation uses the two held-out test sets to assess every model on unseen data. A "
         "reproducibility check rebuilt a fresh virtual environment, ran the test suite and the quick pipeline end "
         "to end (reports/repro_check.md). The conclusions are based on the measured results rather than assumed "
         "values."),
    ]
    for title, text in items:
        heading(doc, title)
        para(doc, text)


def results(doc: Document, F: dict[str, Any]) -> None:
    chapter(doc, 6, "RESULTS AND DISCUSSION", big_title=False)
    para(doc, f"The project, *{TITLE}*, compares manual tuning with AutoML for detecting malicious Windows programs. "
              "The evaluation considers predictive performance on both test sets, tuning methods, time, reliability "
              "and interpretability.")
    rvf, avm = F["rvf"], F["avm"]
    heading(doc, "1. Comparison of Model Performance")
    order = rvf.sort_values("test_future_tpr_at_1pct_fpr", ascending=False).index
    time_of = {"optuna": avm.loc["lgbm_tuned", "wall_minutes"], "lgbm_default": avm.loc["lgbm_default", "wall_minutes"],
               "logreg": avm.loc["logreg", "wall_minutes"], "flaml": avm.loc["flaml", "wall_minutes"],
               "autogluon": avm.loc["autogluon", "wall_minutes"], "h2o": avm.loc["h2o", "wall_minutes"]}
    ts = F["ts"]
    for m in ("grid", "random"):   # search plus refit time, the same definition as the tuned row
        time_of[m] = (ts.loc[m, "search_seconds"] + ts.loc[m, "fit_seconds"]) / 60
    rows = [[MODEL_NAMES[m], f"{rvf.loc[m, 'test_random_roc_auc']:.4f}", f"{rvf.loc[m, 'test_future_roc_auc']:.4f}",
             pct(rvf.loc[m, "test_future_tpr_at_1pct_fpr"]), f"{rvf.loc[m, 'test_future_f1']:.4f}",
             f"{time_of[m]:.1f}"] for m in order]
    para(doc, "Table 6.1 lists every model, ordered by the main metric on the future test set.", after=6)
    data_table(doc, ["Model", "AUC (random test)", "AUC (future test)", "TPR at 1% FPR (future)", "F1 (future)",
                     "Time (min)"], rows, widths=[5.0, 2.2, 2.2, 2.6, 1.8, 1.7])
    caption(doc, "Table 6.1 Results on both test sets (reports/tables/random_vs_future.csv, automl_vs_manual.csv)")
    pw = F["pw"]
    ag = pw["autogluon"]["tpr_at_1pct_fpr"]
    para(doc, f"The best models are tuned LightGBM ({F['tuned_tpr_f']} of future malware caught) and AutoGluon. "
              f"Their difference, {100 * ag.diff_a_minus_b:+.2f} points (95% CI {100 * ag.ci_low:+.2f} to "
              f"{100 * ag.ci_high:+.2f}), {'includes zero, so they are tied' if not ag.significant_95 else 'excludes zero'}. "
              f"FLAML ({100 * pw['flaml']['tpr_at_1pct_fpr'].diff_a_minus_b:+.2f} points) and H2O AutoML "
              f"({100 * pw['h2o']['tpr_at_1pct_fpr'].diff_a_minus_b:+.2f} points) are {behind_phrase(F)}. "
              f"Logistic regression reaches only {F['logreg_auc_f']} AUC on the future test set.")
    heading(doc, "2. Effect of Hyperparameter Optimization")
    para(doc, f"Tuning helped: on the future test set the share of malware caught rose from {F['default_tpr_f']} with "
              f"default settings to {F['tuned_tpr_f']} for the Optuna-tuned model. The search strategy mattered much "
              f"less: Grid, Random and Optuna ended within {F['hand_spread_pts']} points of each other. "
              f"{F['best_cv_method'].capitalize()} had the best cross-validated AUC and was therefore selected, "
              "without looking at any test score.")
    heading(doc, "3. Performance of AutoML Approaches")
    picks = "; ".join(f"{TOOL_NAMES[t]}: {plain_pick(str(avm.loc[t, 'interpretability']))}"
                      for t in ("flaml", "autogluon", "h2o"))
    para(doc, f"AutoGluon matched the tuned model in {F['ag_min']} minutes, against {F['tuned_min']} minutes for the "
              f"selected manual search and {F['all_manual_min']} minutes for all three manual searches. All three "
              f"tools ended up with boosted decision trees, the same kind of model as the manual approach ({picks}). "
              "The time comparison is not like for like: the manual searches had no time cap and searched one model "
              f"family on a {F['tune_subset']}-row subset, while the AutoML budget is a soft limit (FLAML took "
              f"{F['avm'].loc['flaml', 'wall_minutes']:.1f} and H2O AutoML {F['avm'].loc['h2o', 'wall_minutes']:.1f} "
              "minutes in total).")
    heading(doc, "4. Overfitting and Data Leakage Analysis")
    para(doc, f"By AUC, the random split overstates performance only slightly for tree models (gaps of "
              f"{F['tree_gap_min']} to {F['tree_gap_max']}), but much more for logistic regression "
              f"({F['logreg_gap']}). The larger effect is the base rate: the tests are half malware, so "
              f"{F['prec_mix']} of the deployed model's alarms are real, but if only {F['prev_rate']} of files were "
              f"malware about {F['prec_low']} would be, about {F['false_per_true']} false alarms per real one. In the "
              f"leakage experiment, {F['leak_relied']} of {F['leak_models']} models made detection_ratio "
              f"{F['leak_rank_text']}, and their future-test AUCs rose to between {F['leak_auc_min']} and "
              f"{F['leak_auc_max']}. {leak_tools_sentence(F)} The "
              f"first-submission date moved the future-test AUC by at most {F['first_sub_max_delta']}, and removing "
              f"the time-proxy columns changed it by {F['ablate_delta']}.")
    picture(doc, FIGURES / "leakage_jump.png", 14.5)
    caption(doc, "Fig 6.1 Scores with and without the leaky features (reports/figures/leakage_jump.png)")
    heading(doc, "5. Interpretability and Fairness")
    para(doc, f"SHAP shows that the tuned model relies most on the {F['top_families']} feature families. Recall differs "
              f"between slices: at a 1% false-positive threshold the model caught {F['signed_tpr']} of signed malware "
              f"({F['signed_n']} files) against {F['unsigned_tpr']} of unsigned malware. On the {F['chal_n']} "
              f"challenge-set files, malware that antivirus engines first missed, the deployed model caught "
              f"{F['chal_rate']}, against {F['chal_vs']} of ordinary future malware. Balanced class weights changed "
              f"the main metric by {F['cw_delta']} points (95% CI {F['cw_ci']}). The tuned model's probabilities are "
              f"well calibrated: Brier score {F['brier_tuned']}, against {F['brier_logreg']} for logistic "
              "regression.")
    picture(doc, FIGURES / "shap_family_bar.png", 13.0)
    caption(doc, "Fig 6.2 SHAP importance by feature family (reports/figures/shap_family_bar.png)")
    heading(doc, "Discussion")
    para(doc, "The comparison shows that a good AutoML tool can match careful manual tuning on accuracy in less "
              "wall-clock time, while the manual model is simpler to deploy: one file, chosen by cross-validation, "
              "that needs no AutoML library and is explained with TreeSHAP. The evaluation choices mattered more than the choice "
              "between manual tuning and AutoML: testing on later weeks, reporting precision at a realistic base "
              "rate, and checking for leaks changed the picture far more than the differences between the top models.")


def conclusion(doc: Document, F: dict[str, Any]) -> None:
    page_break(doc, gap=38)
    para(doc, "CONCLUSION", align="center", size=14, bold=True)
    pw = F["pw"]["autogluon"]["tpr_at_1pct_fpr"]
    for text in [
        f"The project, *{TITLE}*, evaluated manual hyperparameter tuning and Automated Machine Learning for detecting "
        "malicious Windows programs from static features. Using the EMBER2024 Win64 data, it provides a structured, "
        "reproducible approach to data preparation, tuning, AutoML, critical evaluation and deployment.",
        "The implementation includes LightGBM and logistic-regression baselines, LightGBM tuned with Grid Search, "
        "Random Search and Optuna, and the AutoML frameworks FLAML, AutoGluon and H2O AutoML under an equal time "
        f"budget. Tuned LightGBM caught {F['tuned_tpr_f']} of future malware at a 1% false-positive rate, and "
        f"AutoGluon was statistically tied with it ({100 * pw.diff_a_minus_b:+.2f} points) in {F['ag_min']} minutes.",
        "The project also highlights reliable evaluation practices: a future test set, paired bootstrap confidence "
        f"intervals, leakage experiments, slice metrics and calibration. A planted leak was used as the top input by "
        f"{F['leak_relied']} of {F['leak_models']} models and {F['leak_verdict']}; and at a realistic "
        f"{F['prev_rate']} share of malware most alarms would be false. "
        "SHAP explains the model's decisions, and the FastAPI service and Streamlit page show how the trained model "
        "can be used in an application.",
        "Overall, the project shows that AutoML can match careful manual work on accuracy and saves time, but does "
        "not replace human judgement about how to test a model honestly, how to spot leaks, and where to set the "
        "decision threshold.",
        "In the future, the project can be extended by repeating the AutoML runs with several seeds to measure their "
        "run-to-run variation, testing against adversarially modified malware, evaluating other file types in "
        "EMBER2024, retraining on a sliding window to follow drift, and deploying the application on a cloud "
        "platform with authentication.",
    ]:
        para(doc, text)


REFERENCES: list[tuple[str, str]] = [
    ("Joyce, R. J., Miller, G., Roth, P., Zak, R., Zaresky-Williams, E., Anderson, H., Raff, E., and Holt, J. (2025). "
     "*EMBER2024 - A Benchmark Dataset for Holistic Evaluation of Malware Classifiers.* Proceedings of the 31st ACM "
     "SIGKDD Conference on Knowledge Discovery and Data Mining. Available at: ", "https://arxiv.org/abs/2506.05074"),
    ("Anderson, H. S., and Roth, P. (2018). *EMBER: An Open Dataset for Training Static PE Malware Machine Learning "
     "Models.* arXiv preprint arXiv:1804.04637. Available at: ", "https://arxiv.org/abs/1804.04637"),
    ("Ke, G., Meng, Q., Finley, T., et al. (2017). LightGBM: A Highly Efficient Gradient Boosting Decision Tree. "
     "*Advances in Neural Information Processing Systems, 30.* Available at: ",
     "https://proceedings.neurips.cc/paper/2017/hash/6449f44a102fde848669bdd9eb6b76fa-Abstract.html"),
    ("Akiba, T., Sano, S., Yanase, T., Ohta, T., and Koyama, M. (2019). Optuna: A Next-generation Hyperparameter "
     "Optimization Framework. *Proceedings of the 25th ACM SIGKDD International Conference on Knowledge Discovery & "
     "Data Mining.* Available at: ", "https://doi.org/10.1145/3292500.3330701"),
    ("Wang, C., Wu, Q., Weimer, M., and Zhu, E. (2021). FLAML: A Fast and Lightweight AutoML Library. *Proceedings "
     "of Machine Learning and Systems, 3.* Available at: ", "https://arxiv.org/abs/1911.04706"),
    ("Erickson, N., Mueller, J., Shirkov, A., et al. (2020). AutoGluon-Tabular: Robust and Accurate AutoML for "
     "Structured Data. *arXiv preprint arXiv:2003.06505.* Available at: ", "https://arxiv.org/abs/2003.06505"),
    ("LeDell, E., and Poirier, S. (2020). H2O AutoML: Scalable Automatic Machine Learning. *7th ICML Workshop on "
     "Automated Machine Learning.* Documentation available at: ",
     "https://docs.h2o.ai/h2o/latest-stable/h2o-docs/automl.html"),
    ("Lundberg, S. M., and Lee, S. I. (2017). A Unified Approach to Interpreting Model Predictions. *Advances in "
     "Neural Information Processing Systems, 30.* Available at: ",
     "https://proceedings.neurips.cc/paper/2017/hash/8a20a8621978632d76c43dfd28b67767-Abstract.html"),
    ("Pendlebury, F., Pierazzi, F., Jordaney, R., Kinder, J., and Cavallaro, L. (2019). TESSERACT: Eliminating "
     "Experimental Bias in Malware Classification across Space and Time. *28th USENIX Security Symposium.* "
     "Available at: ", "https://www.usenix.org/conference/usenixsecurity19/presentation/pendlebury"),
    ("Pedregosa, F., Varoquaux, G., Gramfort, A., et al. (2011). Scikit-learn: Machine Learning in Python. *Journal "
     "of Machine Learning Research, 12*, 2825-2830. Available at: ", "https://jmlr.org/papers/v12/pedregosa11a.html"),
    ("Mitchell, M., Wu, S., Zaldivar, A., et al. (2019). Model Cards for Model Reporting. *Proceedings of the "
     "Conference on Fairness, Accountability, and Transparency.* Available at: ", "https://doi.org/10.1145/3287560.3287596"),
    ("FastAPI. *FastAPI Documentation.* Available at: ", "https://fastapi.tiangolo.com/"),
    ("Streamlit. *Streamlit Documentation.* Available at: ", "https://docs.streamlit.io/"),
]


def references(doc: Document) -> None:
    page_break(doc, gap=12)
    para(doc, "REFERENCES", align="left", size=14, bold=True, after=30)
    for i, (text, url) in enumerate(REFERENCES, start=1):
        # left-aligned: justified lines before a long URL would stretch into wide gaps
        p = para(doc, "", align="left", indent=1.27, hanging=0.63)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.27))
        runs(p, f"{i}.\t{text}")
        hyperlink(p, url)


def appendices(doc: Document) -> None:
    page_break(doc, gap=22)
    para(doc, "APPENDICES", align="center", size=14, bold=True, after=12)
    para(doc, "SCREENSHOTS", align="left", size=14, bold=True, after=8)
    for path, width in ((SHOTS / "a_malicious_correct.png", 11.5), (SHOTS / "c_upload_notepad.png", 11.5),
                        (FIGURES / "automl_vs_manual.png", 15.9), (FIGURES / "shap_beeswarm.png", 13.0)):
        picture(doc, path, width, before=8, after=8)


def evaluation_sheet(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    page_number_footer(section)
    para(doc, "EVALUATION SHEET", align="center", size=14, bold=True, before=6, after=6)
    for text in (f"Reg.No :{REG_NO}", f"Name:{STUDENT}", f"Course code: {COURSE_CODE}",
                 f"Course Name: {COURSE_NAME.upper()}"):
        para(doc, text, align="left", bold=True, spacing=1.5)
    para(doc, "", spacing=1.0, after=6)
    rows = [["1", "Online Certification", "10", ""], ["2", "Baseline model and hyperparameter tuning", "5", ""],
            ["3", "AutoML implementation and tool comparison", "10", ""], ["4", "Critical evaluation", "5", ""],
            ["5", "Deployment", "5", ""], ["6", "Documentation and presentation", "5", ""], ["", "Total", "40", ""]]
    t = doc.add_table(rows=len(rows) + 1, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, values in enumerate([["S.No", "Rubrics", "Maximum Marks", "Marks Obtained"]] + rows):
        for c, value in enumerate(values):
            cell = t.cell(r, c)
            cell.width = Cm([1.6, 8.0, 3.2, 3.2][c])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (c == 1 and r > 0) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(4)
            runs(p, value, size=12 if r > 0 and value and c != 2 else 11, bold=(r == 0 or value in ("Total", "40")))
            if r == 0:
                shade(cell, HEADER_FILL)
    para(doc, "Signature of the Faculty-in-charge", align="right", bold=True, before=190)


# --- Build --------------------------------------------------------------------------------------

def build_docx(path: Path = OUT_DOCX) -> Path:
    report_diagrams.main()
    F = facts()
    doc = Document()
    setup(doc)
    cover(doc)
    certificates(doc)
    title_page(doc)
    abstract(doc, F)
    introduction(doc, F)
    literature(doc)
    system_design(doc, F)
    implementation(doc, F)
    testing(doc, F)
    results(doc, F)
    conclusion(doc, F)
    references(doc)
    appendices(doc)
    evaluation_sheet(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.author = STUDENT
    doc.save(path)
    return path


def to_pdf(docx: Path, pdf: Path) -> None:
    """Print the .docx to PDF with Microsoft Word (PowerShell COM), so the PDF matches Word's layout."""
    script = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
              f"$d = $w.Documents.Open('{docx}'); $d.SaveAs([ref]'{pdf}', [ref]17); $d.Close($false); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True, timeout=300)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--no-pdf", action="store_true", help="write the .docx only")
    args = parser.parse_args()
    docx = build_docx()
    print(f"[skill report] wrote {docx.relative_to(ROOT)}")
    if not args.no_pdf:
        to_pdf(docx, OUT_PDF)
        print(f"[skill report] wrote {OUT_PDF.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
